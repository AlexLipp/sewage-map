"""Shared implementation for the eight company EDM event cleaners."""

from __future__ import annotations

import argparse
import csv
from difflib import SequenceMatcher
import hashlib
import importlib
import io
import json
import logging
import math
import re
import shutil
import sys
import time
from collections import Counter
from datetime import date, datetime, time as datetime_time, timedelta
from pathlib import Path
from typing import Any, Iterable, Sequence

import pandas as pd
from openpyxl import load_workbook


TARGET_COLUMNS = [
    "location_name",
    "permit_number",
    "start_time",
    "stop_time",
    "duration_minutes",
]
SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xls"}
COMPANIES = [
    "anglian",
    "northumbria",
    "severn_trent",
    "south_west_water",
    "southern_water",
    "united_utilities",
    "wessex",
    "yorkshire",
]
OUTPUT_FILENAMES: dict[str, str] = {}
TIMESTAMP_FORMAT = "%Y-%m-%dT%H:%M:%S"
TIMESTAMP_PATTERN = re.compile(r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}$")
DURATION_TOLERANCE_MINUTES = 1 / 60
NULL_TEXT = {"", "nan", "none", "nat", "<na>", "null", "n/a", "na"}
LOCATION_PLACEHOLDERS = NULL_TEXT | {"tbc", "unknown", "not available", "not known"}
INTERNAL_COLUMNS = ["_source_file", "_source_sheet", "_source_row_number"]
HEADER_SIMILARITY_THRESHOLD = 0.90
HEADER_SIMILARITY_MARGIN = 0.08
TEMPORAL_PARSE_THRESHOLD = 0.85
TEMPORAL_PAIRED_THRESHOLD = 0.80
TEMPORAL_NONNEGATIVE_THRESHOLD = 0.95

DATETIME_FORMATS = {
    "default": [
        "%Y-%m-%dT%H:%M:%S.%fZ",
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y %H:%M",
    ],
    "anglian": [
        "%Y-%m-%dT%H:%M:%S.%fZ",
        "%Y-%m-%dT%H:%M:%SZ",
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
    ],
}

# Populated by the eight company modules. Keeping configurations there makes
# verified company differences visible without duplicating shared machinery.
COMPANY_SCHEMA_CONFIGS: dict[str, dict[str, Any]] = {}
COMPANY_CLEANERS: dict[str, Any] = {}


def register_company(
    company: str,
    output_filename: str,
    schema_config: dict[str, Any],
    cleaner: Any,
) -> None:
    """Register one company's visible configuration and cleaner entry point."""
    if company not in COMPANIES:
        raise ValueError(f"Unknown company registration: {company}")
    expected_filename = f"{company}_cleaned_data.csv"
    if output_filename != expected_filename:
        raise ValueError(f"Unexpected output filename for {company}: {output_filename}")
    existing = COMPANY_SCHEMA_CONFIGS.get(company)
    if existing is not None and existing != schema_config:
        raise ValueError(f"Conflicting schema registration for {company}")
    COMPANY_SCHEMA_CONFIGS[company] = schema_config
    OUTPUT_FILENAMES[company] = output_filename
    COMPANY_CLEANERS[company] = cleaner


def load_company_modules() -> None:
    """Import all company modules so their configurations are registered."""
    package = "clean_EIR_stopstartdata.raw_to_standardised"
    for company in COMPANIES:
        importlib.import_module(f"{package}.clean_{company}")

RUN_SUMMARY_COLUMNS = [
    "company", "source_file", "source_sheet", "processing_status", "input_rows",
    "structural_rows_excluded", "valid_rows", "retained_rows", "retained_complete_rows",
    "retained_missing_location_rows", "retained_missing_permit_rows",
    "retained_missing_both_rows", "rejected_rows", "output_rows_contributed",
    "exact_duplicate_records", "possible_repeated_source_records", "zero_duration_rows",
    "negative_duration_rows", "start_parse_failures", "stop_parse_failures",
    "location_fallback_rows", "suspicious_permit_rows", "duration_discrepancy_rows",
    "schema_variant", "cleaner_used", "elapsed_seconds", "notes",
]
SCHEMA_INVENTORY_COLUMNS = [
    "company", "source_file", "source_sheet", "input_rows", "original_columns",
    "normalised_columns", "detected_location_column", "detected_permit_column",
    "detected_start_column_or_columns", "detected_stop_column_or_columns",
    "detected_raw_duration_column", "inferred_raw_duration_unit", "schema_variant",
    "cleaner_used", "status", "location_mapping_method", "location_mapping_score",
    "permit_mapping_method", "permit_mapping_score", "start_mapping_method",
    "stop_mapping_method", "missing_location_column", "missing_permit_column", "notes",
]
REJECTED_COLUMNS = [
    "company", "source_file", "source_sheet", "source_row_number",
    "relevant original source values", "rejection_reason",
]
DURATION_COLUMNS = [
    "company", "source_file", "source_sheet", "source_row_number", "location_name",
    "permit_number", "start_time", "stop_time", "raw_duration",
    "inferred_raw_duration_unit", "calculated_duration_minutes",
    "absolute_difference_minutes", "raw_to_calculated_ratio", "discrepancy_category",
]
DUPLICATE_COLUMNS = [
    "company", "source_file", "source_sheet", "source_row_number", "location_name",
    "permit_number", "start_time", "stop_time", "duration_minutes", "duplicate_type",
    "duplicate_group_id",
]
SUSPICIOUS_COLUMNS = [
    "company", "source_file", "source_sheet", "source_row_number", "location_name",
    "permit_number", "reason_flagged",
]

logger = logging.getLogger("standardise_stopstart_data")


class SourceProcessingError(RuntimeError):
    """Raised when a source cannot be interpreted without guessing."""


def find_repository_root(start: Path | None = None) -> Path:
    """Find the repository root from the script location or a supplied path."""
    candidate = (start or Path(__file__)).resolve()
    if candidate.is_file():
        candidate = candidate.parent
    for path in [candidate, *candidate.parents]:
        if (path / ".git").exists() and (path / "raw_data").exists():
            return path
    raise FileNotFoundError("Could not find the SewageMap repository root.")


def resolve_from_root(root: Path, value: str | Path) -> Path:
    """Resolve a CLI path relative to the repository root."""
    path = Path(value)
    return (path if path.is_absolute() else root / path).resolve()


def discover_input_files(input_root: Path, companies: Sequence[str]) -> dict[str, list[Path]]:
    """Recursively discover supported inputs for each selected company."""
    discovered: dict[str, list[Path]] = {}
    for company in companies:
        folder = input_root / company
        discovered[company] = sorted(
            path for path in folder.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_SUFFIXES
        ) if folder.exists() else []
    return discovered


def normalise_column_name(value: Any) -> str:
    """Normalise a header for controlled exact alias matching."""
    text = "" if value is None else str(value)
    text = text.replace("_", " ").replace("\r", " ").replace("\n", " ").lower().strip()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _all_aliases(company: str) -> set[str]:
    aliases: set[str] = set()
    for key, values in COMPANY_SCHEMA_CONFIGS[company].items():
        if key not in {"duration_unit"} and isinstance(values, list):
            aliases.update(normalise_column_name(value) for value in values)
    return aliases


def _unique_headers(values: Sequence[Any]) -> list[str]:
    headers: list[str] = []
    used: Counter[str] = Counter()
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None and str(value).strip() else f"__unnamed_{index}"
        used[base] += 1
        headers.append(base if used[base] == 1 else f"{base}__duplicate_{used[base]}")
    return headers


def _semantic_header_count(headers: Sequence[str], aliases: Sequence[str]) -> int:
    """Count exact headers first; use similarity only when no exact header exists."""
    wanted = [normalise_column_name(alias) for alias in aliases]
    normalised = [normalise_column_name(header) for header in headers]
    exact = [header for header in normalised if header in wanted]
    if exact:
        return len(exact)
    scored = sorted(
        (max(SequenceMatcher(None, header, alias).ratio() for alias in wanted), header)
        for header in normalised if header
    )
    if not scored:
        return 0
    best_score, _ = scored[-1]
    runner_score = scored[-2][0] if len(scored) > 1 else 0.0
    return int(best_score >= HEADER_SIMILARITY_THRESHOLD and best_score - runner_score >= HEADER_SIMILARITY_MARGIN)


def _temporal_header_score(values: Sequence[Any], company: str) -> int:
    """Score only rows containing a complete, recognisable temporal structure."""
    headers = [str(value) for value in values if not _is_blank(value)]
    config = COMPANY_SCHEMA_CONFIGS[company]
    combined = 0
    if config.get("start_datetime"):
        start = _semantic_header_count(headers, config["start_datetime"])
        stop = _semantic_header_count(headers, config["stop_datetime"])
        combined = 2 if start == 1 and stop == 1 else 0
    semantics = ("start_date", "start_time_component", "stop_date", "stop_time_component")
    split = 0
    if all(config.get(key) for key in semantics):
        matches = [_semantic_header_count(headers, config[key]) for key in semantics]
        split = 4 if all(count == 1 for count in matches) else 0
    return max(combined, split)


def _detect_header_row(rows: Iterable[tuple[int, Sequence[Any]]], company: str) -> tuple[int, list[str]]:
    aliases = _all_aliases(company)
    scored: list[tuple[int, int, int, int, list[Any]]] = []
    for row_number, values in rows:
        normalised = [normalise_column_name(value) for value in values]
        temporal_score = _temporal_header_score(values, company)
        alias_score = sum(value in aliases for value in normalised if value)
        nonempty = sum(bool(value) for value in normalised)
        scored.append((temporal_score, alias_score, nonempty, -row_number, list(values)))
    if not scored:
        raise SourceProcessingError("The source contains no rows from which to identify a header.")
    temporal_score, alias_score, _, negative_row, values = max(scored)
    if temporal_score == 0:
        raise SourceProcessingError(
            f"Could not identify a safe temporal header row (best controlled-alias score: {alias_score})."
        )
    return -negative_row, _unique_headers(values)


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    try:
        if pd.isna(value):
            return True
    except (TypeError, ValueError):
        pass
    return isinstance(value, str) and not value.strip()


def _is_repeated_header(values: Sequence[Any], headers: Sequence[str]) -> bool:
    row = [normalise_column_name(value) for value in values[: len(headers)]]
    expected = [normalise_column_name(value) for value in headers]
    return row == expected


def _detect_csv_format(path: Path) -> tuple[str, str]:
    raw = path.read_bytes()[:100_000]
    encoding = "utf-8-sig"
    for candidate in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SourceProcessingError("Could not determine a safe CSV encoding.")
    try:
        delimiter = csv.Sniffer().sniff(text[:20_000], delimiters=",;\t|").delimiter
    except csv.Error:
        delimiter = ","
    return encoding, delimiter


def load_csv_safely(path: Path, company: str) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Load a CSV while preserving identifiers and accounting for structural rows."""
    encoding, delimiter = _detect_csv_format(path)
    parsed: list[tuple[int, list[str]]] = []
    with path.open("r", encoding=encoding, newline="") as handle:
        reader = csv.reader(handle, delimiter=delimiter)
        for row in reader:
            parsed.append((reader.line_num, row))
    header_row, headers = _detect_header_row(parsed[:100], company)
    header_index = next(index for index, (number, _) in enumerate(parsed) if number == header_row)
    records: list[list[Any]] = []
    source_rows: list[int] = []
    structural = 0
    for row_number, row in parsed[header_index + 1 :]:
        padded = list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row))
        if all(_is_blank(value) for value in padded) or _is_repeated_header(padded, headers):
            structural += 1
            continue
        records.append(padded)
        source_rows.append(row_number)
    frame = pd.DataFrame(records, columns=headers)
    frame["_source_row_number"] = source_rows
    return frame, {
        "input_rows": len(parsed) - header_index - 1,
        "structural_rows_excluded": structural,
        "encoding": encoding,
        "delimiter": delimiter,
        "header_row": header_row,
        "columns": headers,
    }


def inspect_excel_workbook(path: Path) -> Any:
    """Open an Excel workbook with the engine appropriate to its extension."""
    if path.suffix.lower() == ".xlsx":
        return load_workbook(path, read_only=True, data_only=True)
    try:
        return pd.ExcelFile(path, engine="xlrd")
    except ImportError as exc:
        raise SourceProcessingError(
            "A legacy .xls file was found. Install xlrd>=2.0, then rerun the pipeline."
        ) from exc


def load_excel_sheet(workbook: Any, sheet: Any, company: str) -> tuple[pd.DataFrame, dict[str, Any]]:
    """Load one event worksheet and retain physical source row numbers."""
    if hasattr(sheet, "iter_rows"):
        preview = [(number, list(row)) for number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 100), values_only=True), start=1
        )]
        header_row, headers = _detect_header_row(preview, company)
        records: list[list[Any]] = []
        source_rows: list[int] = []
        structural = 0
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            values = list(row[: len(headers)]) + [None] * max(0, len(headers) - len(row))
            if all(_is_blank(value) for value in values) or _is_repeated_header(values, headers):
                structural += 1
                continue
            records.append(values)
            source_rows.append(row_number)
        input_rows = max((sheet.max_row or header_row) - header_row, 0)
        frame = pd.DataFrame(records, columns=headers)
        frame["_source_row_number"] = source_rows
        return frame, {
            "input_rows": input_rows,
            "structural_rows_excluded": structural,
            "header_row": header_row,
            "columns": headers,
        }

    sheet_name = str(sheet)
    preview_frame = pd.read_excel(workbook, sheet_name=sheet_name, header=None, nrows=100, dtype=object)
    preview = [(index + 1, row.tolist()) for index, row in preview_frame.iterrows()]
    header_row, _ = _detect_header_row(preview, company)
    frame = pd.read_excel(workbook, sheet_name=sheet_name, header=header_row - 1, dtype=object)
    headers = _unique_headers(list(frame.columns))
    frame.columns = headers
    frame["_source_row_number"] = range(header_row + 1, header_row + 1 + len(frame))
    blank = frame[headers].apply(lambda column: column.map(_is_blank)).all(axis=1)
    repeated = frame[headers].apply(lambda row: _is_repeated_header(row.tolist(), headers), axis=1)
    structural = int((blank | repeated).sum())
    frame = frame.loc[~(blank | repeated)].copy()
    return frame, {
        "input_rows": int(len(frame) + structural),
        "structural_rows_excluded": structural,
        "header_row": header_row,
        "columns": headers,
    }


def _sheet_is_structural(company: str, sheet_name: str, max_columns: int) -> bool:
    name = normalise_column_name(sheet_name)
    if company == "south_west_water":
        return name != "data"
    if company == "wessex":
        return name not in {"discharge time", "discharge times"}
    if company == "united_utilities":
        return "summary" in name or max_columns < 3
    if company == "severn_trent":
        return name in {
            "user guide so", "storm overflows", "drop downs so",
            "edm return stw 2020", "severn trent 2023",
        }
    return False


def _preview_structural_columns(sheet: Any) -> list[str]:
    if not hasattr(sheet, "iter_rows"):
        return []
    candidates: list[list[Any]] = []
    for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, 25), values_only=True):
        candidates.append(list(row))
    if not candidates:
        return []
    values = max(candidates, key=lambda row: sum(not _is_blank(value) for value in row))
    return [str(value).strip() for value in values if not _is_blank(value)]


def _similar_header_candidate(
    columns: Sequence[str], aliases: Sequence[str], semantic: str, frame: pd.DataFrame
) -> tuple[str | None, float | None, str]:
    """Return one unique high-confidence similar header, subject to content checks."""
    scored: list[tuple[float, str]] = []
    for column in columns:
        score = max(
            (SequenceMatcher(None, normalise_column_name(column), normalise_column_name(alias)).ratio()
             for alias in aliases),
            default=0.0,
        )
        scored.append((score, column))
    scored.sort(reverse=True)
    if not scored:
        return None, None, ""
    best_score, best = scored[0]
    runner = scored[1][1] if len(scored) > 1 else ""
    runner_score = scored[1][0] if len(scored) > 1 else 0.0
    if best_score < HEADER_SIMILARITY_THRESHOLD or best_score - runner_score < HEADER_SIMILARITY_MARGIN:
        return None, best_score, runner
    values = clean_text_series(frame[best], location=semantic == "location_name")
    populated = values.dropna()
    if populated.empty:
        return None, best_score, runner
    if semantic == "location_name":
        consistent = float(populated.str.contains(r"[A-Za-z]", regex=True).mean()) >= 0.70
    else:
        consistent = float(populated.str.len().le(80).mean()) >= 0.95
    return (best, best_score, runner) if consistent else (None, best_score, runner)


def _choose_column(
    columns: Sequence[str], aliases: Sequence[str], semantic: str, frame: pd.DataFrame,
    required: bool = True, allow_similarity: bool = False,
) -> tuple[str | None, str, float | None, str]:
    exact = [column for column in columns for alias in aliases if column == alias]
    exact = list(dict.fromkeys(exact))
    if len(exact) == 1:
        return exact[0], "exact_mapping", 1.0, ""
    if len(exact) > 1:
        raise SourceProcessingError(f"Ambiguous {semantic}: exact candidates {exact}.")
    wanted = {normalise_column_name(alias) for alias in aliases}
    normalised = [column for column in columns if normalise_column_name(column) in wanted]
    if len(normalised) == 1:
        return normalised[0], "normalised_exact_mapping", 1.0, ""
    if len(normalised) > 1:
        raise SourceProcessingError(f"Ambiguous {semantic}: normalised candidates {normalised}.")
    if allow_similarity:
        column, score, runner = _similar_header_candidate(columns, aliases, semantic, frame)
        if column:
            return column, "high_confidence_similar_header_mapping", score, runner
    if required:
        raise SourceProcessingError(f"Missing required {semantic}; verified aliases were {list(aliases)}.")
    return None, "no_safe_optional_mapping", None, ""


def _validate_temporal_mapping(frame: pd.DataFrame, mapping: dict[str, Any], company: str) -> None:
    """Require fuzzy temporal mappings to pass paired timestamp content validation."""
    if not any(mapping.get(f"{key}_method") == "high_confidence_similar_header_mapping" for key in (
        "start_datetime", "stop_datetime", "start_date", "start_time_component",
        "stop_date", "stop_time_component",
    )):
        return
    if mapping.get("start_datetime"):
        raw_start = frame[mapping["start_datetime"]]
        raw_stop = frame[mapping["stop_datetime"]]
    else:
        raw_start = combine_date_and_time(frame[mapping["start_date"]], frame[mapping["start_time_component"]])
        raw_stop = combine_date_and_time(frame[mapping["stop_date"]], frame[mapping["stop_time_component"]])
    formats = DATETIME_FORMATS.get(company, DATETIME_FORMATS["default"])
    start = parse_datetime_series(raw_start, formats)
    stop = parse_datetime_series(raw_stop, formats)
    start_nonblank = ~_source_missing(raw_start)
    stop_nonblank = ~_source_missing(raw_stop)
    start_rate = float(start.loc[start_nonblank].notna().mean()) if start_nonblank.any() else 0.0
    stop_rate = float(stop.loc[stop_nonblank].notna().mean()) if stop_nonblank.any() else 0.0
    paired_source = start_nonblank & stop_nonblank
    paired = start.notna() & stop.notna()
    paired_rate = float(paired.sum() / paired_source.sum()) if paired_source.any() else 0.0
    nonnegative_rate = float(stop.loc[paired].ge(start.loc[paired]).mean()) if paired.any() else 0.0
    if (
        start_rate < TEMPORAL_PARSE_THRESHOLD
        or stop_rate < TEMPORAL_PARSE_THRESHOLD
        or paired_rate < TEMPORAL_PAIRED_THRESHOLD
        or nonnegative_rate < TEMPORAL_NONNEGATIVE_THRESHOLD
    ):
        raise SourceProcessingError("Similar temporal headers failed deterministic timestamp content validation.")


def resolve_column_mapping(frame: pd.DataFrame, company: str) -> dict[str, Any]:
    """Resolve a company's required semantic fields using controlled aliases only."""
    config = COMPANY_SCHEMA_CONFIGS[company]
    columns = [column for column in frame.columns if not str(column).startswith("_source_")]
    mapping: dict[str, Any] = {"duration_unit": config.get("duration_unit", "")}
    for key, semantic in (("location", "location_name"), ("permit", "permit_number")):
        column, method, score, runner = _choose_column(
            columns, config[key], semantic, frame, required=False, allow_similarity=True
        )
        mapping[key] = column
        mapping[f"{key}_method"] = method
        mapping[f"{key}_score"] = score
        mapping[f"{key}_competing_candidate"] = runner
    if mapping["location"] and mapping["location"] == mapping["permit"]:
        mapping["location"] = None
        mapping["location_method"] = "no_safe_optional_mapping"
        mapping["location_score"] = None
    if config.get("location_fallback"):
        fallback, method, score, runner = _choose_column(
            columns, config["location_fallback"], "location fallback", frame, required=False
        )
        mapping.update(location_fallback=fallback, location_fallback_method=method,
                       location_fallback_score=score, location_fallback_competing_candidate=runner)
    else:
        mapping["location_fallback"] = None
    if config.get("raw_duration"):
        raw_duration, method, score, runner = _choose_column(
            columns, config["raw_duration"], "raw duration", frame, required=False
        )
        mapping.update(raw_duration=raw_duration, raw_duration_method=method,
                       raw_duration_score=score, raw_duration_competing_candidate=runner)
        if raw_duration and "hrs" in normalise_column_name(raw_duration).split():
            mapping["duration_unit"] = "hours"
    else:
        mapping["raw_duration"] = None
    combined_keys = ("start_datetime", "stop_datetime")
    split_keys = ("start_date", "start_time_component", "stop_date", "stop_time_component")
    combined_mapping: dict[str, Any] = {}
    split_mapping: dict[str, Any] = {}
    if all(config.get(key) for key in combined_keys):
        for key in combined_keys:
            column, method, score, runner = _choose_column(
                columns, config[key], key.replace("_", " "), frame,
                required=False, allow_similarity=True,
            )
            combined_mapping.update({key: column, f"{key}_method": method, f"{key}_score": score,
                                     f"{key}_competing_candidate": runner})
    if all(config.get(key) for key in split_keys):
        for key in split_keys:
            column, method, score, runner = _choose_column(
                columns, config[key], key, frame, required=False, allow_similarity=True,
            )
            split_mapping.update({key: column, f"{key}_method": method, f"{key}_score": score,
                                  f"{key}_competing_candidate": runner})
    combined_complete = bool(combined_mapping) and all(combined_mapping.get(key) for key in combined_keys)
    split_complete = bool(split_mapping) and all(split_mapping.get(key) for key in split_keys)
    if combined_complete and split_complete:
        raise SourceProcessingError("Both combined and split timestamp schemas matched; temporal mapping is ambiguous.")
    if combined_complete:
        mapping.update(combined_mapping)
    elif split_complete:
        mapping.update(split_mapping)
    else:
        raise SourceProcessingError("Failed essential mapping: no complete safe start/stop timestamp schema.")
    temporal_columns = [mapping.get(key) for key in (
        "start_datetime", "stop_datetime", "start_date", "start_time_component",
        "stop_date", "stop_time_component",
    ) if mapping.get(key)]
    if len(temporal_columns) != len(set(temporal_columns)):
        raise SourceProcessingError("Ambiguous timestamp mapping selected one source column more than once.")
    _validate_temporal_mapping(frame, mapping, company)
    return mapping


def clean_text_series(series: pd.Series, *, location: bool = False) -> pd.Series:
    """Trim textual fields, collapse whitespace, and remove literal null markers."""
    nulls = LOCATION_PLACEHOLDERS if location else NULL_TEXT

    def clean(value: Any) -> Any:
        if _is_blank(value):
            return pd.NA
        text = re.sub(r"\s+", " ", str(value).strip())
        return pd.NA if text.casefold() in nulls else text

    return series.map(clean).astype("string")


def clean_identifier_series(series: pd.Series) -> pd.Series:
    """Clean identifiers as text while preserving leading zeroes and punctuation."""
    cleaned = clean_text_series(series)
    return cleaned.str.replace(r"^([+-]?\d+)\.0$", r"\1", regex=True)


def _source_missing(series: pd.Series) -> pd.Series:
    return series.map(lambda value: _is_blank(value) or str(value).strip().casefold() in NULL_TEXT)


def parse_datetime_series(series: pd.Series, formats: Sequence[str]) -> pd.Series:
    """Parse timestamps using only verified explicit formats."""
    result = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    object_mask = series.map(lambda value: isinstance(value, (datetime, date, pd.Timestamp)))
    if object_mask.any():
        result.loc[object_mask] = pd.to_datetime(series.loc[object_mask], errors="coerce")
    text_mask = ~object_mask & ~_source_missing(series)
    unresolved = text_mask.copy()
    text = series.astype("string").str.strip()
    for fmt in formats:
        if not unresolved.any():
            break
        converted = pd.to_datetime(text.loc[unresolved], format=fmt, errors="coerce")
        parsed_mask = converted.notna()
        if parsed_mask.any():
            indices = converted.index[parsed_mask]
            result.loc[indices] = converted.loc[indices]
            unresolved.loc[indices] = False
    return result


def _format_date_component(value: Any) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return pd.Timestamp(value).strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _format_time_component(value: Any) -> str | None:
    if _is_blank(value):
        return None
    if isinstance(value, (datetime, pd.Timestamp)):
        return pd.Timestamp(value).strftime("%H:%M:%S.%f")
    if isinstance(value, datetime_time):
        return value.strftime("%H:%M:%S.%f")
    if isinstance(value, timedelta):
        seconds = value.total_seconds() % 86_400
        return f"{int(seconds // 3600):02d}:{int(seconds % 3600 // 60):02d}:{seconds % 60:09.6f}"
    if isinstance(value, (int, float)) and 0 <= float(value) < 1:
        seconds = float(value) * 86_400
        return f"{int(seconds // 3600):02d}:{int(seconds % 3600 // 60):02d}:{seconds % 60:09.6f}"
    text = str(value).strip()
    for fmt in ("%H:%M:%S.%f", "%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).strftime("%H:%M:%S.%f")
        except ValueError:
            continue
    return None


def combine_date_and_time(date_series: pd.Series, time_series: pd.Series) -> pd.Series:
    """Combine separate Excel date and time fields without losing seconds."""
    values: list[Any] = []
    for date_value, time_value in zip(date_series, time_series):
        rendered_date = _format_date_component(date_value)
        rendered_time = _format_time_component(time_value)
        values.append(f"{rendered_date} {rendered_time}" if rendered_date and rendered_time else pd.NA)
    return pd.Series(values, index=date_series.index, dtype="string")


def calculate_duration_minutes(start: pd.Series, stop: pd.Series) -> pd.Series:
    """Calculate event duration from parsed timestamps, preserving fractional minutes."""
    return ((stop - start).dt.total_seconds() / 60).round(6)


def _repair_verified_source_datetime_coercion(
    series: pd.Series, company: str, context: dict[str, Any]
) -> pd.Series:
    """Repair verified day/month coercion in Anglian's monthly workbooks."""
    filename = Path(context["source_file"]).name.casefold()
    month_names = {
        "january": 1, "february": 2, "march": 3, "april": 4,
        "may": 5, "june": 6, "july": 7, "august": 8,
        "september": 9, "october": 10, "november": 11, "december": 12,
    }
    source_month = next((number for name, number in month_names.items() if name in filename), None)
    if company != "anglian" or source_month is None:
        return series

    def repair(value: Any) -> Any:
        if not isinstance(value, (datetime, pd.Timestamp)):
            return value
        stamp = pd.Timestamp(value)
        # Some monthly releases mix US-style text with Excel-coerced cells. For
        # days 1..12, Excel stored month/day backwards (for example intended
        # 12/2/2025 became 2025-02-12). The filename's publication month,
        # neighbouring uncoerced values, and supplied durations verify the swap.
        if stamp.day == source_month and stamp.month != source_month:
            return stamp.replace(month=source_month, day=stamp.month)
        return value

    return series.map(repair)


def _raw_duration_to_minutes(value: Any, unit: str) -> float | None:
    if _is_blank(value):
        return None
    if unit == "hh:mm:ss":
        if isinstance(value, datetime_time):
            return round((value.hour * 3600 + value.minute * 60 + value.second + value.microsecond / 1_000_000) / 60, 6)
        if isinstance(value, timedelta):
            return round(value.total_seconds() / 60, 6)
        text = str(value).strip()
        match = re.fullmatch(r"(\d+):(\d{1,2}):(\d{1,2}(?:\.\d+)?)", text)
        if not match:
            return None
        return round((int(match.group(1)) * 3600 + int(match.group(2)) * 60 + float(match.group(3))) / 60, 6)
    try:
        number = float(str(value).strip().replace(",", ""))
    except (TypeError, ValueError):
        return None
    if not math.isfinite(number):
        return None
    if unit == "hours":
        return number * 60
    if unit == "seconds":
        return number / 60
    return number


def compare_raw_duration(
    frame: pd.DataFrame,
    mapping: dict[str, Any],
    location: pd.Series,
    permit: pd.Series,
    start: pd.Series,
    stop: pd.Series,
    calculated: pd.Series,
    context: dict[str, Any],
) -> list[dict[str, Any]]:
    """Report source-duration values that differ from timestamp-derived duration."""
    column = mapping.get("raw_duration")
    if not column:
        return []
    unit = mapping.get("duration_unit", "minutes")
    rows: list[dict[str, Any]] = []
    for index in frame.index[start.notna() & stop.notna() & calculated.notna()]:
        raw = frame.at[index, column]
        missing = _is_blank(raw) or str(raw).strip().casefold() in NULL_TEXT
        raw_minutes = None if missing else _raw_duration_to_minutes(raw, unit)
        category = ""
        absolute = math.nan
        ratio = math.nan
        if missing:
            category = "missing_raw_duration"
        elif raw_minutes is None:
            category = "invalid_raw_duration"
        else:
            calc = float(calculated.at[index])
            absolute = abs(raw_minutes - calc)
            ratio = raw_minutes / calc if calc else math.nan
            if absolute <= DURATION_TOLERANCE_MINUTES + 1e-9:
                continue
            raw_number: float | None
            try:
                raw_number = float(str(raw).strip())
            except (TypeError, ValueError):
                raw_number = None
            if absolute <= 1.0:
                category = "rounding_difference"
            elif raw_number is not None and calc > 0 and abs(raw_number / 60 - calc) <= 1.0:
                category = "possible_seconds_instead_of_minutes"
            elif raw_number is not None and calc > 0 and abs(raw_number * 60 - calc) <= 1.0 and unit == "seconds":
                category = "possible_minutes_instead_of_seconds"
            else:
                category = "large_unexplained_difference"
        rows.append({
            "company": context["company"],
            "source_file": context["source_file"],
            "source_sheet": context["source_sheet"],
            "source_row_number": int(frame.at[index, "_source_row_number"]),
            "location_name": location.at[index] if pd.notna(location.at[index]) else "",
            "permit_number": permit.at[index] if pd.notna(permit.at[index]) else "",
            "start_time": start.at[index].strftime(TIMESTAMP_FORMAT),
            "stop_time": stop.at[index].strftime(TIMESTAMP_FORMAT),
            "raw_duration": "" if missing else str(raw),
            "inferred_raw_duration_unit": unit,
            "calculated_duration_minutes": calculated.at[index],
            "absolute_difference_minutes": "" if math.isnan(absolute) else round(absolute, 6),
            "raw_to_calculated_ratio": "" if math.isnan(ratio) else round(ratio, 6),
            "discrepancy_category": category,
        })
    return rows


def _normalised_text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value).strip()).casefold() if pd.notna(value) else ""


def _location_style_identifier(value: Any) -> bool:
    if pd.isna(value):
        return False
    text = str(value).strip()
    words = re.findall(r"[A-Za-z]{3,}", text)
    return len(text) >= 20 and len(words) >= 3 and not re.search(r"\d", text)


def _suspicious_permits(
    frame: pd.DataFrame,
    location: pd.Series,
    permit: pd.Series,
    context: dict[str, Any],
    mapping: dict[str, Any],
) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def add(index: Any, reason: str) -> None:
        rows.append({
            "company": context["company"],
            "source_file": context["source_file"],
            "source_sheet": context["source_sheet"],
            "source_row_number": int(frame.at[index, "_source_row_number"]),
            "location_name": location.at[index] if pd.notna(location.at[index]) else "",
            "permit_number": permit.at[index] if pd.notna(permit.at[index]) else "",
            "reason_flagged": reason,
        })

    for index in frame.index[permit.isna()]:
        add(index, "missing_identifier")
    equal = permit.notna() & location.notna() & permit.map(_normalised_text).eq(location.map(_normalised_text))
    for index in frame.index[equal]:
        add(index, "equals_location_name")
    style = permit.map(_location_style_identifier)
    for index in frame.index[style]:
        add(index, "location_style_identifier")
    if mapping.get("permit"):
        raw = frame[mapping["permit"]]
        type_names = raw.loc[~_source_missing(raw)].map(lambda value: type(value).__name__)
        if type_names.nunique() > 1:
            majority = type_names.value_counts().index[0]
            for index in type_names.index[type_names.ne(majority)]:
                add(index, "unexpected_identifier_type")
    return rows


def _add_reason(reasons: pd.Series, mask: pd.Series, reason: str) -> None:
    current = reasons.loc[mask]
    reasons.loc[mask] = current.map(lambda value: f"{value};{reason}" if value else reason)


def _schema_variant(company: str, mapping: dict[str, Any]) -> str:
    if company == "severn_trent":
        return "severn_trent_ea_and_operational_names" if mapping.get("location_fallback") else "severn_trent_2024"
    if company == "south_west_water":
        return "south_west_water_separate_date_time"
    if company == "united_utilities" and mapping["permit"] == "UUG Reference":
        return "united_utilities_uug_2025"
    if not mapping.get("start_datetime"):
        return f"{company}_separate_date_time"
    return f"{company}_combined_datetime"


def _relevant_values(frame: pd.DataFrame, index: Any, mapping: dict[str, Any]) -> str:
    keys = [
        "location", "location_fallback", "permit", "start_datetime", "start_date",
        "start_time_component", "stop_datetime", "stop_date", "stop_time_component", "raw_duration",
    ]
    values = {
        str(mapping[key]): frame.at[index, mapping[key]]
        for key in keys if mapping.get(key) and mapping[key] in frame.columns
    }
    return json.dumps(values, ensure_ascii=False, default=str)


def _clean_source(frame: pd.DataFrame, company: str, context: dict[str, Any], load_meta: dict[str, Any]) -> dict[str, Any]:
    started = time.perf_counter()
    mapping = resolve_column_mapping(frame, company)
    core_keys = [
        "location", "permit", "start_datetime", "start_date", "start_time_component",
        "stop_datetime", "stop_date", "stop_time_component",
    ]
    core_columns = [mapping[key] for key in core_keys if mapping.get(key)]
    structural_mask = pd.Series(True, index=frame.index)
    for column in core_columns:
        structural_mask &= _source_missing(frame[column])
    added_structural = int(structural_mask.sum())
    if added_structural:
        frame = frame.loc[~structural_mask].copy()
        load_meta = dict(load_meta)
        load_meta["structural_rows_excluded"] += added_structural
    primary_location = (
        clean_text_series(frame[mapping["location"]], location=True)
        if mapping.get("location")
        else pd.Series(pd.NA, index=frame.index, dtype="string")
    )
    location = primary_location.copy()
    fallback_rows = pd.Series(False, index=frame.index)
    if mapping.get("location_fallback"):
        fallback = clean_text_series(frame[mapping["location_fallback"]], location=True)
        fallback_rows = location.isna() & fallback.notna()
        location.loc[fallback_rows] = fallback.loc[fallback_rows]
    permit = (
        clean_identifier_series(frame[mapping["permit"]])
        if mapping.get("permit")
        else pd.Series(pd.NA, index=frame.index, dtype="string")
    )
    formats = DATETIME_FORMATS.get(company, DATETIME_FORMATS["default"])
    if mapping.get("start_datetime"):
        raw_start = _repair_verified_source_datetime_coercion(
            frame[mapping["start_datetime"]], company, context
        )
        raw_stop = _repair_verified_source_datetime_coercion(
            frame[mapping["stop_datetime"]], company, context
        )
    else:
        raw_start = combine_date_and_time(frame[mapping["start_date"]], frame[mapping["start_time_component"]])
        raw_stop = combine_date_and_time(frame[mapping["stop_date"]], frame[mapping["stop_time_component"]])
    start = parse_datetime_series(raw_start, formats)
    stop = parse_datetime_series(raw_stop, formats)
    calculated = calculate_duration_minutes(start, stop)

    missing_start = _source_missing(raw_start)
    missing_stop = _source_missing(raw_stop)
    start_fail = ~missing_start & start.isna()
    stop_fail = ~missing_stop & stop.isna()
    negative = start.notna() & stop.notna() & stop.lt(start)
    nonfinite = calculated.notna() & ~calculated.map(lambda value: math.isfinite(float(value)))
    reasons = pd.Series("", index=frame.index, dtype="string")
    _add_reason(reasons, missing_start, "missing_start_time")
    _add_reason(reasons, missing_stop, "missing_stop_time")
    _add_reason(reasons, start_fail, "unparseable_start_time")
    _add_reason(reasons, stop_fail, "unparseable_stop_time")
    _add_reason(reasons, negative, "negative_duration")
    _add_reason(reasons, nonfinite, "non_finite_duration")
    invalid = reasons.ne("")
    valid = ~invalid

    standardised = pd.DataFrame({
        "location_name": location.loc[valid].astype("string"),
        "permit_number": permit.loc[valid].astype("string"),
        "start_time": start.loc[valid].dt.strftime(TIMESTAMP_FORMAT),
        "stop_time": stop.loc[valid].dt.strftime(TIMESTAMP_FORMAT),
        "duration_minutes": calculated.loc[valid].astype(float),
        "_source_file": context["source_file"],
        "_source_sheet": context["source_sheet"],
        "_source_row_number": frame.loc[valid, "_source_row_number"].astype(int),
    })
    rejected = [{
        "company": company,
        "source_file": context["source_file"],
        "source_sheet": context["source_sheet"],
        "source_row_number": int(frame.at[index, "_source_row_number"]),
        "relevant original source values": _relevant_values(frame, index, mapping),
        "rejection_reason": reasons.at[index],
    } for index in frame.index[invalid]]
    suspicious = _suspicious_permits(frame, location, permit, context, mapping)
    discrepancies = compare_raw_duration(frame, mapping, location, permit, start, stop, calculated, context)
    variant = _schema_variant(company, mapping)
    missing_location = valid & location.isna()
    missing_permit = valid & permit.isna()
    missing_both = missing_location & missing_permit
    complete = valid & location.notna() & permit.notna()
    status = "PARTIAL" if rejected or missing_location.any() or missing_permit.any() else "SUCCESS"
    columns = load_meta["columns"]
    start_columns = [mapping[key] for key in ("start_datetime", "start_date", "start_time_component") if mapping.get(key)]
    stop_columns = [mapping[key] for key in ("stop_datetime", "stop_date", "stop_time_component") if mapping.get(key)]
    notes = []
    if company == "united_utilities" and mapping["permit"] == "UUG Reference":
        notes.append("Source supplies UUG Reference only; preserved without inventing a UUP identifier.")
    if fallback_rows.any():
        notes.append(f"Used verified location fallback for {int(fallback_rows.sum())} row(s).")
    if not mapping.get("location"):
        notes.append("No safe optional location column mapping; retained events with blank location_name.")
    if not mapping.get("permit"):
        notes.append("No safe optional permit column mapping; retained events with blank permit_number.")
    summary = {
        "company": company,
        "source_file": context["source_file"],
        "source_sheet": context["source_sheet"],
        "processing_status": status,
        "input_rows": load_meta["input_rows"],
        "structural_rows_excluded": load_meta["structural_rows_excluded"],
        "valid_rows": int(valid.sum()),
        "retained_rows": int(valid.sum()),
        "retained_complete_rows": int(complete.sum()),
        "retained_missing_location_rows": int(missing_location.sum()),
        "retained_missing_permit_rows": int(missing_permit.sum()),
        "retained_missing_both_rows": int(missing_both.sum()),
        "rejected_rows": int(invalid.sum()),
        "output_rows_contributed": int(valid.sum()),
        "exact_duplicate_records": 0,
        "possible_repeated_source_records": 0,
        "zero_duration_rows": int((valid & calculated.eq(0)).sum()),
        "negative_duration_rows": int(negative.sum()),
        "start_parse_failures": int(start_fail.sum()),
        "stop_parse_failures": int(stop_fail.sum()),
        "location_fallback_rows": int(fallback_rows.sum()),
        "suspicious_permit_rows": len(suspicious),
        "duration_discrepancy_rows": len(discrepancies),
        "schema_variant": variant,
        "cleaner_used": f"clean_{company}",
        "elapsed_seconds": round(time.perf_counter() - started, 3),
        "notes": " ".join(notes),
    }
    inventory = {
        "company": company,
        "source_file": context["source_file"],
        "source_sheet": context["source_sheet"],
        "input_rows": load_meta["input_rows"],
        "original_columns": json.dumps(columns, ensure_ascii=False),
        "normalised_columns": json.dumps([normalise_column_name(column) for column in columns], ensure_ascii=False),
        "detected_location_column": (mapping.get("location") or "") + (f"; fallback={mapping['location_fallback']}" if mapping.get("location_fallback") else ""),
        "detected_permit_column": mapping.get("permit") or "",
        "detected_start_column_or_columns": json.dumps(start_columns, ensure_ascii=False),
        "detected_stop_column_or_columns": json.dumps(stop_columns, ensure_ascii=False),
        "detected_raw_duration_column": mapping.get("raw_duration") or "",
        "inferred_raw_duration_unit": mapping.get("duration_unit") or "",
        "schema_variant": variant,
        "cleaner_used": f"clean_{company}",
        "status": status,
        "location_mapping_method": mapping.get("location_method", "no_safe_optional_mapping"),
        "location_mapping_score": mapping.get("location_score") if mapping.get("location_score") is not None else "",
        "permit_mapping_method": mapping.get("permit_method", "no_safe_optional_mapping"),
        "permit_mapping_score": mapping.get("permit_score") if mapping.get("permit_score") is not None else "",
        "start_mapping_method": mapping.get("start_datetime_method") or ";".join(
            str(mapping.get(f"{key}_method", "")) for key in ("start_date", "start_time_component")
        ),
        "stop_mapping_method": mapping.get("stop_datetime_method") or ";".join(
            str(mapping.get(f"{key}_method", "")) for key in ("stop_date", "stop_time_component")
        ),
        "missing_location_column": not bool(mapping.get("location")),
        "missing_permit_column": not bool(mapping.get("permit")),
        "notes": " ".join(notes),
    }
    return {
        "data": standardised,
        "summary": summary,
        "inventory": inventory,
        "rejected": rejected,
        "suspicious": suspicious,
        "duration_discrepancies": discrepancies,
    }


def clean_company_source(
    frame: pd.DataFrame,
    company: str,
    context: dict[str, Any],
    load_meta: dict[str, Any],
) -> dict[str, Any]:
    """Apply the registered company's verified schema to one source table."""
    return _clean_source(frame, company, context, load_meta)


def _stable_group_id(prefix: str, values: Sequence[Any]) -> str:
    payload = json.dumps([str(value) for value in values], ensure_ascii=False, separators=(",", ":"))
    return f"{prefix}-{hashlib.sha1(payload.encode('utf-8')).hexdigest()[:12]}"


def classify_duplicate_events(company: str, frame: pd.DataFrame) -> list[dict[str, Any]]:
    """Identify exact records and identity collisions without grouping normal sensor events."""
    diagnostics: list[dict[str, Any]] = []
    exact_mask = frame.duplicated(TARGET_COLUMNS, keep=False)
    for index in frame.index[exact_mask]:
        row = frame.loc[index]
        diagnostics.append({
            "company": company,
            "source_file": row["_source_file"],
            "source_sheet": row["_source_sheet"],
            "source_row_number": int(row["_source_row_number"]),
            **{column: row[column] for column in TARGET_COLUMNS},
            "duplicate_type": "exact_duplicate_record",
            "duplicate_group_id": _stable_group_id("exact", [row[column] for column in TARGET_COLUMNS]),
        })
    identity = ["permit_number", "start_time", "stop_time"]
    identity_mask = frame.duplicated(identity, keep=False)
    candidates = frame.loc[identity_mask].copy()
    if not candidates.empty:
        candidates["_variant_key"] = (
            candidates["location_name"].astype(str)
            + "\x1f"
            + candidates["duration_minutes"].map(lambda value: f"{float(value):.6f}")
        )
        variant_count = candidates.groupby(identity, sort=False, dropna=False)["_variant_key"].transform("nunique")
        for _, row in candidates.loc[variant_count.gt(1)].iterrows():
            group_id = _stable_group_id("possible", [row[column] for column in identity])
            diagnostics.append({
                "company": company,
                "source_file": row["_source_file"],
                "source_sheet": row["_source_sheet"],
                "source_row_number": int(row["_source_row_number"]),
                **{column: row[column] for column in TARGET_COLUMNS},
                "duplicate_type": "possible_repeated_source_record",
                "duplicate_group_id": group_id,
            })
    return diagnostics


def validate_standardised_dataframe(frame: pd.DataFrame) -> None:
    """Validate event fields while allowing missing optional identity metadata."""
    if list(frame.columns) != TARGET_COLUMNS:
        raise ValueError(f"Output columns are not exact: {list(frame.columns)}")
    if any(str(column).lower().startswith("unnamed") for column in frame.columns):
        raise ValueError("An accidental index column is present.")
    if frame.empty:
        raise ValueError("The consolidated output is empty.")
    for column in ("location_name", "permit_number"):
        populated = frame.loc[frame[column].notna() & frame[column].astype("string").str.strip().ne(""), column]
        if not populated.map(lambda value: isinstance(value, str)).all():
            raise ValueError(f"{column} contains a populated non-string value.")
        if populated.astype("string").str.strip().str.casefold().isin(NULL_TEXT - {""}).any():
            raise ValueError(f"{column} contains literal null text.")
    for column in ("start_time", "stop_time"):
        if not frame[column].map(lambda value: bool(TIMESTAMP_PATTERN.fullmatch(str(value)))).all():
            raise ValueError(f"{column} does not use exact YYYY-MM-DDTHH:MM:SS formatting.")
    start = pd.to_datetime(frame["start_time"], format=TIMESTAMP_FORMAT, errors="coerce")
    stop = pd.to_datetime(frame["stop_time"], format=TIMESTAMP_FORMAT, errors="coerce")
    if start.isna().any() or stop.isna().any():
        raise ValueError("A formatted timestamp cannot be reparsed.")
    if stop.lt(start).any():
        raise ValueError("An output event has stop_time earlier than start_time.")
    duration = pd.to_numeric(frame["duration_minutes"], errors="coerce")
    if duration.isna().any() or not duration.map(math.isfinite).all() or duration.lt(0).any():
        raise ValueError("duration_minutes contains an invalid value.")
    expected = (stop - start).dt.total_seconds() / 60
    if (duration.sub(expected).abs() > DURATION_TOLERANCE_MINUTES + 1e-9).any():
        raise ValueError("duration_minutes does not match the formatted timestamps within one second.")


def _empty_row(columns: Sequence[str], **values: Any) -> dict[str, Any]:
    return {column: values.get(column, "") for column in columns}


def _record_structural_sheet(
    state: dict[str, Any], company: str, source_file: str, sheet_name: str,
    input_rows: int, columns: list[str], elapsed: float,
) -> None:
    notes = "Non-event cover, readme, summary, guidance, or metadata sheet excluded."
    state["run_summary"].append(_empty_row(
        RUN_SUMMARY_COLUMNS,
        company=company, source_file=source_file, source_sheet=sheet_name,
        processing_status="SUCCESS", input_rows=input_rows,
        structural_rows_excluded=input_rows, valid_rows=0, rejected_rows=0,
        output_rows_contributed=0, exact_duplicate_records=0,
        possible_repeated_source_records=0, zero_duration_rows=0,
        negative_duration_rows=0, start_parse_failures=0, stop_parse_failures=0,
        location_fallback_rows=0, suspicious_permit_rows=0,
        duration_discrepancy_rows=0, schema_variant="structural_sheet",
        cleaner_used=f"clean_{company}", elapsed_seconds=round(elapsed, 3), notes=notes,
    ))
    state["schema_inventory"].append(_empty_row(
        SCHEMA_INVENTORY_COLUMNS,
        company=company, source_file=source_file, source_sheet=sheet_name,
        input_rows=input_rows, original_columns=json.dumps(columns, ensure_ascii=False),
        normalised_columns=json.dumps([normalise_column_name(value) for value in columns], ensure_ascii=False),
        schema_variant="structural_sheet", cleaner_used=f"clean_{company}",
        status="SUCCESS", notes=notes,
    ))


def _failure_block(
    company: str, source_file: str, source_sheet: str, exc: Exception,
    columns: Sequence[str], attempted: dict[str, Any] | None = None,
) -> str:
    return "\n".join([
        f"company: {company}",
        f"relative source path: {source_file}",
        f"source sheet: {source_sheet}",
        f"exception type: {type(exc).__name__}",
        f"message: {exc}",
        f"original columns: {json.dumps(list(columns), ensure_ascii=False)}",
        f"normalised columns: {json.dumps([normalise_column_name(value) for value in columns], ensure_ascii=False)}",
        f"attempted semantic mappings: {json.dumps(attempted or COMPANY_SCHEMA_CONFIGS[company], ensure_ascii=False)}",
        "reason safe processing was impossible: required semantics were missing or ambiguous; no guess was made.",
        "suggested next action: inspect the source schema and add a verified exact alias or explicit schema variant.",
        "-" * 72,
    ])


def _record_failed_source(
    state: dict[str, Any], company: str, source_file: str, source_sheet: str,
    exc: Exception, columns: Sequence[str], input_rows: int, structural: int,
    frame: pd.DataFrame | None, elapsed: float,
) -> None:
    data_rows = max(input_rows - structural, 0)
    state["failed_files"].append(_failure_block(company, source_file, source_sheet, exc, columns))
    state["run_summary"].append(_empty_row(
        RUN_SUMMARY_COLUMNS,
        company=company, source_file=source_file, source_sheet=source_sheet,
        processing_status="FAILED", input_rows=input_rows,
        structural_rows_excluded=structural, valid_rows=0, rejected_rows=data_rows,
        output_rows_contributed=0, exact_duplicate_records=0,
        possible_repeated_source_records=0, zero_duration_rows=0,
        negative_duration_rows=0, start_parse_failures=0, stop_parse_failures=0,
        location_fallback_rows=0, suspicious_permit_rows=0,
        duration_discrepancy_rows=0, schema_variant="unsupported_schema",
        cleaner_used=f"clean_{company}", elapsed_seconds=round(elapsed, 3), notes=str(exc),
    ))
    state["schema_inventory"].append(_empty_row(
        SCHEMA_INVENTORY_COLUMNS,
        company=company, source_file=source_file, source_sheet=source_sheet,
        input_rows=input_rows, original_columns=json.dumps(list(columns), ensure_ascii=False),
        normalised_columns=json.dumps([normalise_column_name(value) for value in columns], ensure_ascii=False),
        schema_variant="unsupported_schema", cleaner_used=f"clean_{company}",
        status="FAILED", notes=str(exc),
    ))
    if frame is not None:
        raw_columns = [column for column in frame.columns if not str(column).startswith("_source_")]
        for index, row in frame.iterrows():
            state["rejected_rows"].append({
                "company": company,
                "source_file": source_file,
                "source_sheet": source_sheet,
                "source_row_number": int(row.get("_source_row_number", index + 1)),
                "relevant original source values": json.dumps(
                    {column: row[column] for column in raw_columns}, ensure_ascii=False, default=str
                ),
                "rejection_reason": "unsupported_schema",
            })


def _consume_result(state: dict[str, Any], result: dict[str, Any]) -> None:
    state["run_summary"].append(result["summary"])
    state["schema_inventory"].append(result["inventory"])
    state["rejected_rows"].extend(result["rejected"])
    state["suspicious_permit_numbers"].extend(result["suspicious"])
    state["duration_discrepancies"].extend(result["duration_discrepancies"])


def consolidate_company_data(
    root: Path, company: str, files: Sequence[Path], state: dict[str, Any]
) -> pd.DataFrame:
    """Process every current source for one company and concatenate valid events."""
    frames: list[pd.DataFrame] = []
    cleaner = COMPANY_CLEANERS[company]
    for path in files:
        relative = path.relative_to(root).as_posix()
        logger.info("  %s", relative)
        if path.suffix.lower() == ".csv":
            started = time.perf_counter()
            frame: pd.DataFrame | None = None
            metadata: dict[str, Any] = {"input_rows": 0, "structural_rows_excluded": 0, "columns": []}
            try:
                frame, metadata = load_csv_safely(path, company)
                context = {"company": company, "source_file": relative, "source_sheet": ""}
                result = cleaner(frame, context, metadata)
                _consume_result(state, result)
                frames.append(result["data"])
            except (OSError, UnicodeError, csv.Error, ValueError, TypeError, SourceProcessingError) as exc:
                _record_failed_source(
                    state, company, relative, "", exc, metadata.get("columns", []),
                    metadata.get("input_rows", 0), metadata.get("structural_rows_excluded", 0),
                    frame, time.perf_counter() - started,
                )
            continue

        try:
            workbook = inspect_excel_workbook(path)
        except (OSError, ValueError, SourceProcessingError) as exc:
            _record_failed_source(state, company, relative, "", exc, [], 0, 0, None, 0)
            continue
        try:
            if hasattr(workbook, "worksheets"):
                sheets = workbook.worksheets
            else:
                sheets = workbook.sheet_names
            for sheet in sheets:
                started = time.perf_counter()
                sheet_name = sheet.title if hasattr(sheet, "title") else str(sheet)
                max_columns = sheet.max_column if hasattr(sheet, "max_column") else int(
                    pd.read_excel(workbook, sheet_name=sheet_name, header=None, nrows=5).shape[1]
                )
                if _sheet_is_structural(company, sheet_name, max_columns):
                    input_rows = int(sheet.max_row if hasattr(sheet, "max_row") else len(
                        pd.read_excel(workbook, sheet_name=sheet_name, header=None)
                    ))
                    _record_structural_sheet(
                        state, company, relative, sheet_name, input_rows,
                        _preview_structural_columns(sheet), time.perf_counter() - started,
                    )
                    continue
                frame = None
                metadata = {"input_rows": 0, "structural_rows_excluded": 0, "columns": []}
                try:
                    frame, metadata = load_excel_sheet(workbook, sheet, company)
                    context = {"company": company, "source_file": relative, "source_sheet": sheet_name}
                    result = cleaner(frame, context, metadata)
                    _consume_result(state, result)
                    frames.append(result["data"])
                except (OSError, ValueError, TypeError, KeyError, SourceProcessingError) as exc:
                    _record_failed_source(
                        state, company, relative, sheet_name, exc, metadata.get("columns", []),
                        metadata.get("input_rows", 0), metadata.get("structural_rows_excluded", 0),
                        frame, time.perf_counter() - started,
                    )
        finally:
            workbook.close()
    if not frames:
        return pd.DataFrame(columns=TARGET_COLUMNS + INTERNAL_COLUMNS)
    return pd.concat(frames, ignore_index=True)


def _update_duplicate_counts(state: dict[str, Any], diagnostics: list[dict[str, Any]]) -> None:
    exact = Counter()
    possible = Counter()
    for row in diagnostics:
        key = (row["company"], row["source_file"], row["source_sheet"])
        if row["duplicate_type"] == "exact_duplicate_record":
            exact[key] += 1
        else:
            possible[key] += 1
    for row in state["run_summary"]:
        key = (row["company"], row["source_file"], row["source_sheet"])
        row["exact_duplicate_records"] = exact[key]
        row["possible_repeated_source_records"] = possible[key]


def _assert_safe_output(path: Path, root: Path) -> None:
    resolved = path.resolve()
    protected = {
        root.resolve(),
        (root / "raw_data").resolve(),
        (root / "clean_EIR_stopstartdata").resolve(),
        Path.home().resolve(),
    }
    if resolved in protected or not resolved.is_relative_to(root.resolve()):
        raise ValueError(f"Refusing unsafe output path: {resolved}")


def safely_replace_output(staging: Path, output_root: Path, root: Path) -> None:
    """Atomically replace generated output after validating that paths are safe."""
    _assert_safe_output(output_root, root)
    _assert_safe_output(staging, root)
    backup = output_root.parent / f".{output_root.name}.backup-{os_process_id()}"
    _assert_safe_output(backup, root)
    if backup.exists():
        shutil.rmtree(backup)
    try:
        if output_root.exists():
            output_root.rename(backup)
        staging.rename(output_root)
    except Exception:
        if not output_root.exists() and backup.exists():
            backup.rename(output_root)
        raise
    if backup.exists():
        shutil.rmtree(backup)


def os_process_id() -> int:
    """Return the current process ID without adding platform-specific path logic."""
    import os
    return os.getpid()


def _initial_state() -> dict[str, Any]:
    return {
        "run_summary": [],
        "schema_inventory": [],
        "rejected_rows": [],
        "failed_files": [],
        "duration_discrepancies": [],
        "duplicate_events": [],
        "suspicious_permit_numbers": [],
        "output_rows": {},
        "final_validation_failures": [],
    }


def _file_status_counts(rows: list[dict[str, Any]]) -> Counter[str]:
    priority = {"SUCCESS": 0, "PARTIAL": 1, "FAILED": 2}
    per_file: dict[str, str] = {}
    for row in rows:
        path = row["source_file"]
        status = row["processing_status"]
        if path not in per_file or priority[status] > priority[per_file[path]]:
            per_file[path] = status
    return Counter(per_file.values())


def _print_summary(root: Path, input_root: Path, output_root: Path, discovered: dict[str, list[Path]], state: dict[str, Any]) -> None:
    counts = _file_status_counts(state["run_summary"])
    summaries = state["run_summary"]
    logger.info("\n=== Stop/start standardisation summary ===")
    logger.info("Repository root: %s", root)
    logger.info("Raw input root: %s", input_root)
    logger.info("Output root: %s", output_root)
    logger.info("Files discovered: %s", sum(len(paths) for paths in discovered.values()))
    logger.info("Files succeeded: %s", counts["SUCCESS"])
    logger.info("Files partially succeeded: %s", counts["PARTIAL"])
    logger.info("Files failed: %s", counts["FAILED"])
    logger.info("Total loaded rows: %s", sum(int(row["input_rows"]) for row in summaries))
    logger.info("Total valid rows: %s", sum(int(row["valid_rows"]) for row in summaries))
    logger.info("Total rejected rows: %s", sum(int(row["rejected_rows"]) for row in summaries))
    logger.info("Exact duplicate records found: %s", sum(row["duplicate_type"] == "exact_duplicate_record" for row in state["duplicate_events"]))
    logger.info("Possible repeated source records found: %s", sum(row["duplicate_type"] == "possible_repeated_source_record" for row in state["duplicate_events"]))
    logger.info("Suspicious permit rows: %s", len(state["suspicious_permit_numbers"]))
    logger.info("Duration discrepancy rows: %s", len(state["duration_discrepancies"]))
    for company, count in state["output_rows"].items():
        logger.info("Output rows - %s: %s", company, count)
    logger.info("Repeated permits with different event timestamps are normal and were not counted as duplicates.")


def run_pipeline(args: argparse.Namespace) -> int:
    """Execute discovery, cleaning, validation, staging, diagnostics, and replacement."""
    root = find_repository_root()
    input_root = resolve_from_root(root, args.input_root)
    output_root = resolve_from_root(root, args.output_root)
    _assert_safe_output(output_root, root)
    selected = [args.company] if args.company else COMPANIES
    discovered = discover_input_files(input_root, selected)
    state = _initial_state()
    staging = output_root.parent / f".{output_root.name}.staging-{os_process_id()}"
    _assert_safe_output(staging, root)
    if staging.exists():
        shutil.rmtree(staging)
    staging.mkdir(parents=True)

    try:
        for company in selected:
            logger.info("\nProcessing %s (%s file(s))", company, len(discovered[company]))
            if not discovered[company]:
                exc = SourceProcessingError(f"No supported files found in {input_root / company}")
                _record_failed_source(state, company, str((input_root / company).relative_to(root)), "", exc, [], 0, 0, None, 0)
                state["output_rows"][company] = 0
                continue
            combined = consolidate_company_data(root, company, discovered[company], state)
            if combined.empty:
                message = "No valid output rows were produced."
                state["final_validation_failures"].append(f"{company}: {message}")
                state["output_rows"][company] = 0
                logger.error("%s: %s", company, message)
                continue
            combined = combined.sort_values(
                ["start_time", "stop_time", "permit_number", "location_name"], kind="mergesort"
            ).reset_index(drop=True)
            duplicates = classify_duplicate_events(company, combined)
            state["duplicate_events"].extend(duplicates)
            _update_duplicate_counts(state, state["duplicate_events"])
            final = combined[TARGET_COLUMNS].copy()
            try:
                validate_standardised_dataframe(final)
            except ValueError as exc:
                state["final_validation_failures"].append(f"{company}: {exc}")
                state["output_rows"][company] = 0
                logger.error("Final validation failed for %s: %s", company, exc)
                continue
            state["output_rows"][company] = len(final)
            if not args.dry_run:
                final.to_csv(
                    staging / OUTPUT_FILENAMES[company], index=False, encoding="utf-8", na_rep=""
                )
            del combined, final

        if args.dry_run:
            shutil.rmtree(staging)
        elif len(selected) == 1:
            company = selected[0]
            staged_output = staging / OUTPUT_FILENAMES[company]
            if staged_output.exists():
                output_root.mkdir(parents=True, exist_ok=True)
                final_output = output_root / OUTPUT_FILENAMES[company]
                _assert_safe_output(final_output, root)
                staged_output.replace(final_output)
            shutil.rmtree(staging)
        else:
            safely_replace_output(staging, output_root, root)
    except Exception:
        if staging.exists():
            shutil.rmtree(staging)
        raise

    _print_summary(root, input_root, output_root, discovered, state)
    failed = any(row["processing_status"] == "FAILED" for row in state["run_summary"])
    exit_failure = failed or bool(state["final_validation_failures"])
    if args.strict:
        exit_failure = exit_failure or any(row["rejected_rows"] for row in state["run_summary"])
        exit_failure = exit_failure or any(state["output_rows"].get(company, 0) == 0 for company in selected)
        exit_failure = exit_failure or bool(state["suspicious_permit_numbers"])
    return 1 if exit_failure else 0


def run_self_checks() -> None:
    """Run in-memory checks for identity, temporal, mapping, and duplicate rules."""
    starts = pd.Series(pd.to_datetime([
        "2025-01-01 10:00:00", "2025-01-01 10:00:00", "2025-01-01 23:59:30",
        "2025-01-01 00:00:00", "2025-01-01 12:00:00",
    ]))
    stops = pd.Series(pd.to_datetime([
        "2025-01-01 10:00:30", "2025-01-01 10:01:30", "2025-01-02 00:00:30",
        "2025-01-04 00:00:00", "2025-01-01 12:00:00",
    ]))
    duration = calculate_duration_minutes(starts, stops)
    assert duration.tolist() == [0.5, 1.5, 1.0, 4320.0, 0.0]
    formatted = pd.DataFrame({
        "location_name": ["Example"] * 5,
        "permit_number": ["00123"] * 5,
        "start_time": starts.dt.strftime(TIMESTAMP_FORMAT),
        "stop_time": stops.dt.strftime(TIMESTAMP_FORMAT),
        "duration_minutes": duration,
    })
    assert list(formatted.columns) == TARGET_COLUMNS
    validate_standardised_dataframe(formatted)
    incomplete = pd.DataFrame({
        "location_name": pd.Series([pd.NA, "Example", pd.NA], dtype="string"),
        "permit_number": pd.Series(["P1", pd.NA, pd.NA], dtype="string"),
        "start_time": ["2025-01-01T10:00:00", "2025-01-02T10:00:00", "2025-01-03T10:00:00"],
        "stop_time": ["2025-01-01T11:00:00", "2025-01-02T11:00:00", "2025-01-03T11:00:00"],
        "duration_minutes": [60.0, 60.0, 60.0],
    })
    validate_standardised_dataframe(incomplete)
    buffer = io.StringIO()
    incomplete.to_csv(buffer, index=False, na_rep="")
    rendered = buffer.getvalue()
    assert "nan" not in rendered.casefold() and "none" not in rendered.casefold() and "<na>" not in rendered.casefold()
    assert ",P1,2025" in rendered and "Example,,2025" in rendered and "\n,,2025" in rendered
    assert clean_identifier_series(pd.Series(["00123", "00123.0"])).tolist() == ["00123", "00123"]
    suspicious_frame = pd.DataFrame({"permit": ["Example Site"], "_source_row_number": [2]})
    suspicious = _suspicious_permits(
        suspicious_frame, pd.Series(["Example Site"]), pd.Series(["Example Site"]),
        {"company": "anglian", "source_file": "memory", "source_sheet": ""}, {"permit": "permit"},
    )
    assert any(row["reason_flagged"] == "equals_location_name" for row in suspicious)
    negative = pd.DataFrame({
        "location_name": ["Example"], "permit_number": ["00123"],
        "start_time": ["2025-01-02T00:00:00"], "stop_time": ["2025-01-01T00:00:00"],
        "duration_minutes": [-1440.0],
    })
    try:
        validate_standardised_dataframe(negative)
    except ValueError:
        pass
    else:
        raise AssertionError("Negative duration was not rejected.")
    context = {"company": "anglian", "source_file": "memory", "source_sheet": ""}
    metadata = {
        "input_rows": 5, "structural_rows_excluded": 0,
        "columns": ["SITE NAME", "SITE CODE", "START DATE/TIME", "END DATE/TIME"],
    }
    source = pd.DataFrame({
        "SITE NAME": [pd.NA, "B", pd.NA, "D", "E"],
        "SITE CODE": ["P1", pd.NA, pd.NA, "P4", "P5"],
        "START DATE/TIME": ["2025-01-01 10:00", "2025-01-02 10:00", "2025-01-03 10:00", "bad", "2025-01-05 11:00"],
        "END DATE/TIME": ["2025-01-01 11:00", "2025-01-02 11:00", "2025-01-03 11:00", "2025-01-04 11:00", "2025-01-05 10:00"],
        "_source_row_number": [2, 3, 4, 5, 6],
    })
    cleaned = _clean_source(source, "anglian", context, metadata)
    assert len(cleaned["data"]) == 3
    assert cleaned["summary"]["retained_missing_location_rows"] == 2
    assert cleaned["summary"]["retained_missing_permit_rows"] == 2
    assert cleaned["summary"]["retained_missing_both_rows"] == 1
    assert cleaned["summary"]["start_parse_failures"] == 1
    assert cleaned["summary"]["negative_duration_rows"] == 1
    assert len(cleaned["rejected"]) == 2

    similar_optional = pd.DataFrame({
        "SITE NME": ["Example"], "SITE CODE": ["P1"],
        "START DATE/TIME": ["2025-01-01 10:00"], "END DATE/TIME": ["2025-01-01 11:00"],
        "_source_row_number": [2],
    })
    similar_mapping = resolve_column_mapping(similar_optional, "anglian")
    assert similar_mapping["location"] == "SITE NME"
    assert similar_mapping["location_method"] == "high_confidence_similar_header_mapping"

    ambiguous_optional = similar_optional.rename(columns={"SITE NME": "SITE NAM"}).copy()
    ambiguous_optional["SITE NAMEE"] = "Example Two"
    ambiguous_mapping = resolve_column_mapping(ambiguous_optional, "anglian")
    assert ambiguous_mapping["location"] is None

    ambiguous_temporal = similar_optional.rename(columns={"SITE NME": "SITE NAME"}).copy()
    ambiguous_temporal["START DATE TIME"] = ambiguous_temporal["START DATE/TIME"]
    try:
        resolve_column_mapping(ambiguous_temporal, "anglian")
    except SourceProcessingError:
        pass
    else:
        raise AssertionError("Ambiguous timestamp mapping did not fail safely.")
    duplicate_input = pd.DataFrame([
        ["A", "P1", "2025-01-01T00:00:00", "2025-01-01T00:30:00", 30.0, "one", "", 2],
        ["A", "P1", "2025-01-02T00:00:00", "2025-01-02T00:30:00", 30.0, "one", "", 3],
        ["A", "P1", "2025-01-01T00:00:00", "2025-01-01T00:30:00", 30.0, "two", "", 2],
    ], columns=TARGET_COLUMNS + INTERNAL_COLUMNS)
    duplicate_rows = classify_duplicate_events("anglian", duplicate_input)
    exact = [row for row in duplicate_rows if row["duplicate_type"] == "exact_duplicate_record"]
    possible = [row for row in duplicate_rows if row["duplicate_type"] == "possible_repeated_source_record"]
    assert len(exact) == 2 and not possible
    assert len(duplicate_input) == 3
    logger.info("All built-in self-checks passed.")


def build_parser(fixed_company: str | None = None) -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    if fixed_company is None:
        parser.add_argument("--company", choices=COMPANIES, help="Process one company only.")
    parser.add_argument("--dry-run", action="store_true", help="Inspect and validate without replacing final CSVs.")
    parser.add_argument("--strict", action="store_true", help="Fail on rejected rows, empty outputs, or suspicious permits.")
    parser.add_argument("--self-check", action="store_true", help="Run in-memory checks and exit.")
    parser.add_argument("--input-root", default="raw_data", help="Raw input root, relative to the repository root.")
    parser.add_argument(
        "--output-root", default="clean_EIR_stopstartdata/input_stopstart_data",
        help="Generated output root, relative to the repository root.",
    )
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return parser


def main(fixed_company: str | None = None) -> int:
    load_company_modules()
    args = build_parser(fixed_company).parse_args()
    if fixed_company is not None:
        args.company = fixed_company
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s: %(message)s")
    if args.self_check:
        run_self_checks()
        return 0
    try:
        return run_pipeline(args)
    except (OSError, ValueError, SourceProcessingError) as exc:
        logger.error("Pipeline stopped safely: %s", exc)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
